"""ATS-friendly DOCX and PDF exports for EggyPDF Career Pro."""
from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape


SECTION_HEADINGS = {
    "summary", "professional summary", "profile", "professional profile", "objective",
    "experience", "work experience", "professional experience", "employment", "work history",
    "education", "academic background", "qualifications",
    "skills", "technical skills", "core competencies", "key skills", "competencies",
    "projects", "certifications", "certificates", "languages", "awards", "achievements",
}


def _lines(text: str):
    return [x.strip() for x in (text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n") if x.strip()]


def _clean_heading(text: str):
    return re.sub(r"[^a-z ]+", " ", (text or "").lower()).strip()


def _is_heading(line: str):
    cleaned = _clean_heading(line)
    if cleaned in SECTION_HEADINGS:
        return True
    words = line.split()
    return bool(1 <= len(words) <= 4 and line.isupper() and len(line) <= 45)


def _is_bullet(line: str):
    return bool(re.match(r"^[•●▪◦*\-]\s+", line))


def _strip_bullet(line: str):
    return re.sub(r"^[•●▪◦*\-]\s+", "", line).strip()


def export_docx(resume_text: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    lines = _lines(resume_text)
    if not lines:
        raise ValueError("Optimized resume text is required.")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    for i, line in enumerate(lines):
        if i == 0 and not _is_heading(line) and not _is_bullet(line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            continue

        if _is_heading(line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line.upper())
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(11.5)
            continue

        if _is_bullet(line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.first_line_indent = Inches(-0.1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(_strip_bullet(line))
            run.font.name = "Arial"
            run.font.size = Pt(10.5)
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def export_pdf(resume_text: str) -> bytes:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    lines = _lines(resume_text)
    if not lines:
        raise ValueError("Optimized resume text is required.")

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="EggyPDF Optimized Resume",
    )
    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "ResumeBody", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5,
        leading=12, spaceAfter=3,
    )
    name_style = ParagraphStyle(
        "ResumeName", parent=body, fontName="Helvetica-Bold", fontSize=16,
        leading=19, alignment=TA_CENTER, spaceAfter=5,
    )
    heading = ParagraphStyle(
        "ResumeHeading", parent=body, fontName="Helvetica-Bold", fontSize=10.5,
        leading=13, spaceBefore=7, spaceAfter=3,
    )
    bullet = ParagraphStyle(
        "ResumeBullet", parent=body, leftIndent=10, firstLineIndent=-7,
        bulletIndent=2, spaceAfter=2,
    )

    story = []
    for i, line in enumerate(lines):
        safe = escape(line)
        if i == 0 and not _is_heading(line) and not _is_bullet(line):
            story.append(Paragraph(safe, name_style))
        elif _is_heading(line):
            story.append(Paragraph(escape(line.upper()), heading))
        elif _is_bullet(line):
            story.append(Paragraph("• " + escape(_strip_bullet(line)), bullet))
        else:
            story.append(Paragraph(safe, body))
    story.append(Spacer(1, 2))
    doc.build(story)
    return out.getvalue()


def build_resume_export(resume_text: str, file_format: str):
    fmt = (file_format or "").strip().lower()
    if fmt == "docx":
        return export_docx(resume_text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "EggyPDF-Optimized-Resume.docx"
    if fmt == "pdf":
        return export_pdf(resume_text), "application/pdf", "EggyPDF-Optimized-Resume.pdf"
    raise ValueError("Export format must be docx or pdf.")
